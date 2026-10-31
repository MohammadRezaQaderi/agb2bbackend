from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_export_workbook(
    rows: list[dict[str, Any]],
    headers: list[tuple[str, str]] | None = None,
) -> bytes:
    """Build a Persian, right-to-left XLSX workbook for an authenticated export."""
    if headers is None:
        headers = [("نام", "first_name"), ("نام خانوادگی", "last_name"), ("مشاور", "con_name")]
        if rows and "access_state" in rows[0]:
            headers.extend([("ذره‌بین", "ag_state"), ("دوپامین", "scl_state")])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "گزارش"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"
    sheet.append([label for label, _ in headers])

    for row in rows:
        output = []
        for _, key in headers:
            if key in ("ag_state", "scl_state"):
                product = "AG" if key == "ag_state" else "SCL"
                access = (row.get("access_state") or {}).get(product) or {}
                if access.get("permission") != 1:
                    value = "بدون دسترسی"
                elif access.get("state") == "completed":
                    value = "تکمیل شده"
                elif access.get("state") == "in-progress":
                    value = f"در حال انجام - {access.get('current_quiz_name') or '-'}"
                else:
                    value = "شروع نشده"
            else:
                value = row.get(key)
            output.append(value)
        sheet.append(output)

    header_fill = PatternFill("solid", fgColor="1676B8")
    for cell in sheet[1]:
        cell.font = Font(name="Tahoma", bold=True, color="FFFFFF", size=12)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            cell.font = Font(name="Tahoma", size=11)
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 3, 15), 50)
    if sheet.max_row > 1 and sheet.max_column > 0:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"

    summary = workbook.create_sheet("خلاصه")
    summary.sheet_view.rightToLeft = True
    summary.append(["خلاصه گزارش", ""])
    summary.append(["تعداد رکوردها", len(rows)])
    summary.append(["تعداد فیلدها", len(headers)])
    for cell in summary[1]:
        cell.font = Font(name="Tahoma", bold=True, color="FFFFFF")
        cell.fill = header_fill
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 24

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
